"""Reads the Excel pool and writes data/entries.json, teams.json, standings.json.
Exports only names, picks and tiebreakers (no emails/addresses).
Usage: python scripts/build_entries.py pool.xlsm"""
import sys, re, json, difflib
from openpyxl import load_workbook
T = {  # abbr: (name, conference, division)
"BUF":("Buffalo","AFC","East"),"MIA":("Miami","AFC","East"),"NE":("New England","AFC","East"),"NYJ":("NY Jets","AFC","East"),
"BAL":("Baltimore","AFC","North"),"CIN":("Cincinnati","AFC","North"),"CLE":("Cleveland","AFC","North"),"PIT":("Pittsburgh","AFC","North"),
"HOU":("Houston","AFC","South"),"IND":("Indianapolis","AFC","South"),"JAX":("Jacksonville","AFC","South"),"TEN":("Tennessee","AFC","South"),
"DEN":("Denver","AFC","West"),"KC":("Kansas City","AFC","West"),"LV":("Las Vegas","AFC","West"),"LAC":("LA Chargers","AFC","West"),
"DAL":("Dallas","NFC","East"),"NYG":("NY Giants","NFC","East"),"PHI":("Philadelphia","NFC","East"),"WSH":("Washington","NFC","East"),
"CHI":("Chicago","NFC","North"),"DET":("Detroit","NFC","North"),"GB":("Green Bay","NFC","North"),"MIN":("Minnesota","NFC","North"),
"ATL":("Atlanta","NFC","South"),"CAR":("Carolina","NFC","South"),"NO":("New Orleans","NFC","South"),"TB":("Tampa Bay","NFC","South"),
"ARI":("Arizona","NFC","West"),"LAR":("LA Rams","NFC","West"),"SF":("San Francisco","NFC","West"),"SEA":("Seattle","NFC","West")}
byname = {v[0].lower(): k for k, v in T.items()}
alias = {"nola": "NO", "tenn": "TEN", "cleve": "CLE"}
wb = load_workbook(sys.argv[1], data_only=True)

price = {}  # team values from the Rules sheet
for row in wb["Rules 2026-2027"].iter_rows(values_only=True):
    c = [x for x in row if x is not None]
    if len(c) > 1 and isinstance(c[0], (int, float)) and all(isinstance(x, str) for x in c[1:]):
        for n in c[1:]: price[byname[n.lower()]] = c[0]
assert len(price) == 32, f"found {len(price)} team prices"

def lookup(s):
    k = s.strip().strip(".").lower()
    if k in byname: return byname[k], False
    if k in alias: return alias[k], False
    m = difflib.get_close_matches(k, byname, 1, 0.7)
    return (byname[m[0]], True) if m else (None, True)

ws = wb["Summary"]; entries = []
for r in range(3, ws.max_row + 1):
    name, teams, tb = ws.cell(r, 2).value, ws.cell(r, 4).value, ws.cell(r, 5).value
    if not name or not teams: continue
    abbrs = []
    for part in re.split(r"[;,]", teams):
        if not part.strip(): continue
        a, fuzzy = lookup(part)
        if a is None: print("UNMATCHED:", name, "->", part); continue
        if fuzzy: print("check:", name, repr(part.strip()), "->", T[a][0])
        abbrs.append(a)
    val = sum(price[a] for a in abbrs)
    if val > 25: print(f"OVER BUDGET: {name} ${val}")
    entries.append({"n": str(name).strip(), "t": abbrs, "tb": tb})

st = {}  # seed current records from the Summary standings block
for r in range(3, 50):
    a = byname.get(str(ws.cell(r, 6).value or "").lower())
    if a: st[a] = {"w": ws.cell(r, 7).value or 0, "l": ws.cell(r, 8).value or 0, "div": False, "pw": 0, "conf": False, "sb": False}
json.dump(entries, open("data/entries.json", "w"))
json.dump({a: {"name": v[0], "conf": v[1], "div": v[2], "price": price[a]} for a, v in T.items()}, open("data/teams.json", "w"))
json.dump({"updated": None, "teams": st}, open("data/standings.json", "w"))
print(len(entries), "entries;", len(st), "teams seeded")
